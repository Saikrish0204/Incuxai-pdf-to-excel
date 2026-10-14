import io
import re
from datetime import datetime

import fitz  # PyMuPDF
import pandas as pd
import pdfplumber
import pytesseract
import streamlit as st
from PIL import Image
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Incuxai | PDF to Excel AI Agent", page_icon="📄", layout="wide")

st.markdown("""
<style>
.main {background-color: #f7faff;}
.hero {padding: 1.2rem 1.4rem; border-radius: 18px; background: linear-gradient(110deg,#09265d,#1555a6); color: white;}
.hero h1 {color:white; margin-bottom:0.25rem;}
.hero p {color:#dbeafe; margin-bottom:0;}
.small-note {color:#53627a; font-size:0.9rem;}
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="hero"><h1>📄 Incuxai Document Intelligence</h1><p>Convert PDF documents into structured Excel workbooks — extract, inspect, validate, export.</p></div>', unsafe_allow_html=True)
st.write("")
with st.expander("How it works"):
    st.write("1. Upload a PDF (maximum 50 pages). 2. The app tries table extraction on each page. 3. If a page has no extractable text/table, OCR is attempted. 4. Review the results and download an Excel workbook.")
st.caption("Prototype notice: OCR and table extraction can make mistakes. Always review important values against the original PDF.")

def clean_cell(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()

def normalise_table(raw):
    rows = []
    for row in raw or []:
        cleaned = [clean_cell(x) for x in row]
        if any(cleaned):
            rows.append(cleaned)
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # Treat the first row as headers only if it contains useful labels.
    first = rows[0]
    nonempty = [x for x in first if x]
    if nonempty and len(set(nonempty)) == len(nonempty):
        headers = []
        seen = {}
        for i, cell in enumerate(first):
            label = cell or f"Column_{i+1}"
            seen[label] = seen.get(label, 0) + 1
            headers.append(label if seen[label] == 1 else f"{label}_{seen[label]}")
        data = rows[1:] if len(rows) > 1 else []
    else:
        headers = [f"Column_{i+1}" for i in range(width)]
        data = rows
    return pd.DataFrame(data, columns=headers)

def ocr_page(page, page_num):
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    text = pytesseract.image_to_string(image, config="--psm 6")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return None, "OCR found no readable text"
    # OCR engines sometimes preserve table columns with multiple spaces/tabs.
    split_lines = [re.split(r"\s{2,}|\t+", line) for line in lines]
    width = max(len(line) for line in split_lines)
    if width >= 2:
        padded = [line + [""] * (width - len(line)) for line in split_lines]
        df = normalise_table(padded)
        return df, f"OCR fallback used on page {page_num}"
    return pd.DataFrame({"Extracted text": lines}), f"OCR text fallback used on page {page_num}; table columns may need manual review"

def make_excel(dataframes, logs):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if not dataframes:
            pd.DataFrame({"Message": ["No tables or readable text were extracted."]}).to_excel(writer, index=False, sheet_name="Extracted Data")
        else:
            for idx, (page_num, df) in enumerate(dataframes, start=1):
                safe_name = f"Page{page_num}_Table{idx}"[:31]
                df.to_excel(writer, index=False, sheet_name=safe_name)
        pd.DataFrame(logs).to_excel(writer, index=False, sheet_name="Extraction Log")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1455A0")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                max_len = max((len(str(c.value or "")) for c in list(col)[:100]), default=10)
                ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 36)
    output.seek(0)
    return output.getvalue()

uploaded = st.file_uploader("Choose a PDF file", type=["pdf"], help="Maximum 50 pages. Scanned documents use OCR if Tesseract is available.")
if uploaded:
    pdf_bytes = uploaded.getvalue()
    if len(pdf_bytes) > 25 * 1024 * 1024:
        st.error("Please upload a PDF smaller than 25 MB.")
        st.stop()
    try:
        pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page_count = len(pdf_doc)
        if page_count == 0:
            st.error("This PDF has no pages.")
            st.stop()
        if page_count > 50:
            st.error(f"This file has {page_count} pages. The challenge limit is 50 pages.")
            st.stop()
        st.info(f"File: **{uploaded.name}** · Pages: **{page_count}** · Size: **{len(pdf_bytes)/1024:.1f} KB**")
    except Exception as exc:
        st.error(f"Could not open this PDF: {exc}")
        st.stop()

    if st.button("🚀 Extract tables and text", type="primary", use_container_width=True):
        extracted = []
        logs = []
        progress = st.progress(0)
        status = st.empty()
        try:
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                for i, page in enumerate(pdf.pages, start=1):
                    status.write(f"Processing page {i} of {page_count}…")
                    page_frames = []
                    try:
                        tables = page.extract_tables()
                        for table in tables or []:
                            df = normalise_table(table)
                            if df is not None and not df.empty:
                                page_frames.append(df)
                        if page_frames:
                            for df in page_frames:
                                extracted.append((i, df))
                            logs.append({"Page": i, "Method": "PDF table extraction", "Status": "Success", "Rows extracted": sum(len(df) for df in page_frames), "Notes": f"{len(page_frames)} table(s) detected"})
                        else:
                            # OCR is used when the PDF page has little/no extractable text.
                            text = page.extract_text() or ""
                            if len(text.strip()) >= 20:
                                lines = [line.strip() for line in text.splitlines() if line.strip()]
                                df = pd.DataFrame({"Extracted text": lines})
                                extracted.append((i, df))
                                logs.append({"Page": i, "Method": "PDF text extraction", "Status": "Review", "Rows extracted": len(df), "Notes": "No table detected; text is saved line-by-line"})
                            else:
                                try:
                                    df, note = ocr_page(pdf_doc[i-1], i)
                                    if df is not None and not df.empty:
                                        extracted.append((i, df))
                                        logs.append({"Page": i, "Method": "OCR", "Status": "Review", "Rows extracted": len(df), "Notes": note})
                                    else:
                                        logs.append({"Page": i, "Method": "OCR", "Status": "Needs review", "Rows extracted": 0, "Notes": note})
                                except Exception as ocr_exc:
                                    logs.append({"Page": i, "Method": "OCR", "Status": "Failed", "Rows extracted": 0, "Notes": f"OCR error: {ocr_exc}. Check Tesseract installation."})
                    except Exception as page_exc:
                        logs.append({"Page": i, "Method": "Extraction", "Status": "Failed", "Rows extracted": 0, "Notes": str(page_exc)})
                    progress.progress(i / page_count)
            status.success("Processing finished. Review the extracted data before downloading.")
            st.session_state["incuxai_extracted"] = extracted
            st.session_state["incuxai_logs"] = logs
            st.session_state["incuxai_filename"] = uploaded.name
        except Exception as exc:
            st.error(f"Processing failed: {exc}")

if "incuxai_extracted" in st.session_state:
    extracted = st.session_state["incuxai_extracted"]
    logs = st.session_state["incuxai_logs"]
    st.subheader("Extraction results")
    c1, c2, c3 = st.columns(3)
    c1.metric("Pages processed", len(logs))
    c2.metric("Tables / data blocks", len(extracted))
    c3.metric("Rows extracted", sum(len(df) for _, df in extracted))
    if extracted:
        for idx, (page_num, df) in enumerate(extracted, start=1):
            with st.expander(f"Page {page_num} · Data block {idx}", expanded=(idx == 1)):
                st.dataframe(df, use_container_width=True, hide_index=True)
    else:
        st.warning("No data could be extracted. Try a clearer PDF or check the OCR installation.")
    st.subheader("Validation and extraction log")
    log_df = pd.DataFrame(logs)
    st.dataframe(log_df, use_container_width=True, hide_index=True)
    warning_count = sum(1 for item in logs if item.get("Status") != "Success")
    st.write(f"**Review flags:** {warning_count} page(s) need review or had extraction issues.")
    excel_bytes = make_excel(extracted, logs)
    st.download_button(
        "⬇️ Download Excel workbook",
        data=excel_bytes,
        file_name=f"{st.session_state.get('incuxai_filename','converted').rsplit('.',1)[0]}_converted.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
    st.download_button(
        "⬇️ Download extraction log as CSV",
        data=log_df.to_csv(index=False).encode("utf-8"),
        file_name="extraction_validation_log.csv",
        mime="text/csv",
        use_container_width=True,
    )

st.divider()
st.caption("Incuxai AI Agent Development Challenge · Student prototype · Built with Python, Streamlit, PDFPlumber, PyMuPDF, Tesseract OCR and Pandas")
